"""Fill the presentation template with your own training results.

Run after training and evaluating both models (Step 5):

    python presentation/make_presentation.py

It reads reports/metrics_*.json and models/improved.keras, renders the
activation maps and the real misclassified tiles, and writes
presentation/EuroSAT_CNN_presentation.pptx.
"""
from __future__ import annotations

import json
import os
import re
import sys
from pathlib import Path

os.environ.setdefault("TF_CPP_MIN_LOG_LEVEL", "2")
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import keras  # noqa: E402
import matplotlib.cm as cm  # noqa: E402
import numpy as np  # noqa: E402
from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402

from eurosat_cnn.config import CLASSES, DATA_DIR, MODELS_DIR, REPORTS_DIR  # noqa: E402
from eurosat_cnn.data import list_images, make_dataset, stratified_split  # noqa: E402
from eurosat_cnn.models import RandomRot90  # noqa: E402,F401

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "presentation_template.pptx"
OUTPUT = HERE / "EuroSAT_CNN_presentation.pptx"
ASSETS = HERE / "build_assets"
PRETTY = {"AnnualCrop": "Annual crop", "Forest": "Forest", "HerbaceousVegetation": "Herbaceous veg.",
          "Highway": "Highway", "Industrial": "Industrial", "Pasture": "Pasture",
          "PermanentCrop": "Permanent crop", "Residential": "Residential", "River": "River", "SeaLake": "Sea / lake"}


def load_rgb(path) -> np.ndarray:
    return np.asarray(Image.open(path).convert("RGB"))


def save_up(arr: np.ndarray, name: str, resample=Image.LANCZOS):
    Image.fromarray(arr).resize((384, 384), resample).save(ASSETS / name)


def activation_image(act, name: str):
    """Save the most varied channel of a layer's output as a coloured heat map."""
    a = np.asarray(act)[0]
    channel = int(np.argmax(a.reshape(-1, a.shape[-1]).std(axis=0)))
    m = a[..., channel]
    m = (m - m.min()) / (m.max() - m.min() + 1e-8)
    save_up((cm.viridis(m)[..., :3] * 255).astype(np.uint8), name, Image.BICUBIC)


def collect() -> dict:
    for m in ("baseline", "improved"):
        if not (REPORTS_DIR / f"metrics_{m}.json").exists():
            sys.exit(f"Missing reports/metrics_{m}.json - run `python -m eurosat_cnn.evaluate` first.")
    metrics = {m: json.loads((REPORTS_DIR / f"metrics_{m}.json").read_text()) for m in ("baseline", "improved")}
    model = keras.models.load_model(MODELS_DIR / "improved.keras")
    ASSETS.mkdir(exist_ok=True)

    paths, labels = list_images(DATA_DIR)
    _, _, test = stratified_split(paths, labels)
    probs = model.predict(make_dataset(test, batch_size=256), verbose=0)
    pred = probs.argmax(1)

    # 1) the residential test tile the model is most sure about, with two activation maps
    res = CLASSES.index("Residential")
    i = max(np.where(test.labels == res)[0], key=lambda j: probs[j, res])
    tile = load_rgb(test.paths[i]).astype("float32")[None] / 255
    save_up((tile[0] * 255).astype(np.uint8), "input.png")
    relus = [layer for layer in model.layers if isinstance(layer, keras.layers.ReLU)]
    probe = keras.Model(model.inputs, [relus[0].output, relus[min(4, len(relus) - 1)].output])
    early, deep = probe(tile, training=False)
    activation_image(early, "fmap1.png")
    activation_image(deep, "fmap2.png")
    top = int(probs[i].argmax())

    # 2) four confident mistakes, one for each of the most frequent confusions
    wrong = np.where(pred != test.labels)[0]
    mistakes = []
    for pair in metrics["improved"]["top_confusions"]:
        t, p = CLASSES.index(pair["true"]), CLASSES.index(pair["predicted"])
        cand = [j for j in wrong if test.labels[j] == t and pred[j] == p]
        if cand:
            j = max(cand, key=lambda k: probs[k, p])
            mistakes.append(j)
    for j in wrong[np.argsort(-probs[wrong].max(1))]:       # top up if fewer than 4 pairs
        if len(mistakes) >= 4:
            break
        if j not in mistakes:
            mistakes.append(j)
    values = {}
    for n, j in enumerate(mistakes[:4]):
        save_up(load_rgb(test.paths[j]), f"mistake{n}.png")
        values[f"M{n}_TRUE"] = PRETTY[CLASSES[test.labels[j]]]
        values[f"M{n}_PRED"] = PRETTY[CLASSES[pred[j]]]
        values[f"M{n}_PROB"] = f"{probs[j].max():.0%}"

    b, im = metrics["baseline"], metrics["improved"]
    err_b, err_i = 1 - b["test_accuracy"], 1 - im["test_accuracy"]
    fewer = (1 - err_i / err_b) if err_b > 0 else 0
    change = "fewer" if fewer >= 0 else "more"
    tc = im["top_confusions"] + [{"true": "-", "predicted": "-", "count": 0}] * 2
    pr = lambda c: PRETTY.get(c, c).lower()  # noqa: E731
    values.update({
        "BASELINE_ACC": f"{b['test_accuracy']:.1%}",
        "IMPROVED_ACC": f"{im['test_accuracy']:.1%}",
        "DECISION_LABEL": PRETTY[CLASSES[top]],
        "DECISION_LABEL_LOWER": PRETTY[CLASSES[top]].lower(),
        "DECISION_PROB": f"{probs[i, top]:.0%}",
        "RESULTS_SENTENCE": (
            f"Out of {im['test_images']:,} unseen images, the improved model labels "
            f"{round(im['test_accuracy'] * im['test_images']):,} correctly: "
            f"{abs(fewer):.0%} {change} mistakes than the baseline."),
        "MISTAKES_SENTENCE": (
            f"The most common mix-ups are {pr(tc[0]['true'])} taken for {pr(tc[0]['predicted'])} "
            f"({tc[0]['count']} images) and {pr(tc[1]['true'])} taken for {pr(tc[1]['predicted'])} "
            f"({tc[1]['count']}). These classes look alike from space, and some tiles contain both."),
    })
    return {"values": values, "metrics": metrics}


def fill(data: dict):
    prs = Presentation(TEMPLATE)
    values = data["values"]
    token = re.compile(r"\{\{(\w+)\}\}")
    sub = lambda text: token.sub(lambda m: values.get(m.group(1), m.group(0)), text)  # noqa: E731

    for slide in prs.slides:
        for shape in list(slide.shapes):
            name = shape.name or ""
            if name.startswith("IMGLABEL:"):
                shape._element.getparent().remove(shape._element)
                continue
            if name.startswith("IMG:"):
                file = ASSETS / name[4:]
                slide.shapes.add_picture(str(file), shape.left, shape.top, shape.width, shape.height)
                shape._element.getparent().remove(shape._element)
                continue
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        if "{{" in run.text:
                            run.text = sub(run.text)
            if getattr(shape, "has_chart", False) and shape.has_chart:
                chart = shape.chart
                if chart.has_title and "F1" in chart.chart_title.text_frame.text:
                    cd = CategoryChartData()
                    cd.categories = [PRETTY[c] for c in CLASSES]
                    for m in ("baseline", "improved"):
                        f1 = data["metrics"][m]["per_class_f1"]
                        cd.add_series(m.capitalize(), [round(f1[c] * 100, 1) for c in CLASSES])
                    chart.replace_data(cd)
        if slide.has_notes_slide:
            for para in slide.notes_slide.notes_text_frame.paragraphs:
                for run in para.runs:
                    run.text = sub(run.text)

    prs.save(OUTPUT)
    print(f"Presentation written to {OUTPUT}")


if __name__ == "__main__":
    fill(collect())
